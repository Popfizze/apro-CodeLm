from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from .models import Finding
from .rebar import norm_text

GT_COLUMNS = ("FEUILLET", "LOCALISATION", "PLAN", "DESSIN")


def _compact(text: str) -> str:
    return norm_text(text).replace(" ", "")


def _gt_row(values: tuple, columns: list[int]) -> dict | None:
    sheet, location, plan, shop = (values[i] for i in columns)
    if sheet is None:
        return None
    plan, shop = str(plan or "").strip(), str(shop or "").strip()
    return {
        "feuillet": str(sheet).strip(),
        "localisation": str(location or "").strip(),
        "plan": plan,
        "atelier": shop,
        "ecart_attendu": _compact(plan) != _compact(shop),
    }


def load_gt(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(workbook[workbook.sheetnames[0]].iter_rows(values_only=True))
    workbook.close()
    header = [norm_text(str(h or "")) for h in rows[0]]
    columns = [next(i for i, h in enumerate(header) if h.startswith(name)) for name in GT_COLUMNS]
    parsed = (_gt_row(values, columns) for values in rows[1:] if values)
    return [row for row in parsed if row is not None]


def _location_key(text: str) -> str:
    return re.sub(r"\s+", "", norm_text(text))


def _candidates(row: dict, findings: list[Finding]) -> list[Finding]:
    location = _location_key(row["localisation"])
    return [
        f for f in findings if f.feuillet == row["feuillet"] and _location_key(f.element or "") == location
    ]


def _reference_fields(reference: Finding | None) -> dict:
    if reference is None:
        return {"finding_id": None, "statut": None}
    return {"finding_id": reference.id, "statut": reference.statut}


def _detail(row: dict, findings: list[Finding]) -> tuple[dict, Finding | None]:
    candidates = _candidates(row, findings)
    hit = next((f for f in candidates if f.statut == "non_conforme"), None)
    reference = hit or next(iter(candidates), None)
    expected = row["ecart_attendu"]
    detail = {
        **row,
        "trouve": bool(hit) if expected else None,
        **_reference_fields(reference),
        "ecarts": [e.model_dump() for e in hit.ecarts] if hit else [],
    }
    return detail, hit if expected else None


def _ratio(numerator: int, denominator: int) -> float | None:
    return round(numerator / denominator, 3) if denominator else None


def _scores(details: list[dict], findings: list[Finding], matched_ids: set[str]) -> dict:
    expected = [d for d in details if d["ecart_attendu"]]
    found = [d for d in expected if d["trouve"]]
    flagged = [f for f in findings if f.statut == "non_conforme"]
    true_positives = [f for f in flagged if f.id in matched_ids]
    return {
        "rappel": _ratio(len(found), len(expected)),
        "precision": _ratio(len(true_positives), len(flagged)),
        "n_attendus": len(expected),
        "n_trouves": len(found),
        "n_non_conformes": len(flagged),
        "hors_verite_terrain": len(flagged) - len(true_positives),
    }


def evaluate(findings: list[Finding], gt_path: Path) -> dict:
    details, matched_ids = [], set()
    for row in load_gt(gt_path):
        detail, matched = _detail(row, findings)
        details.append(detail)
        if matched is not None:
            matched_ids.add(matched.id)
    return {"verite_terrain": gt_path.name, **_scores(details, findings, matched_ids), "details": details}
